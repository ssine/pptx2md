"""Regression tests for pptx2md.multi_column."""

from pptx import Presentation
from pptx.util import Inches

from pptx2md.multi_column import is_two_column_text

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _add_positionless_textbox(slide, text):
    """Add a textbox whose XML has no <a:off>, so shape.top/left are None.

    Real .pptx files can contain shapes that omit an explicit offset.
    """
    tb = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1))
    tb.text_frame.text = text
    xfrm = tb._element.spPr.find(_A + "xfrm")
    xfrm.remove(xfrm.find(_A + "off"))
    return tb


def test_is_two_column_text_handles_positionless_shape():
    # A slide mixing positioned shapes with a shape that has no offset (top/left
    # is None) must not crash while sorting or computing centroids.
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    for i in range(2):
        tb = slide.shapes.add_textbox(Inches(1 + i * 4), Inches(1), Inches(3), Inches(2))
        tb.text_frame.text = f"column {i}"
    _add_positionless_textbox(slide, "no offset")

    mu, sigma = is_two_column_text(slide)

    # The two positioned shapes contribute; the position-less one is skipped.
    assert len(mu) == 2
    assert len(sigma) == 2
